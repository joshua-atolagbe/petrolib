'''
Writing pay summaries to formatted Excel workbooks

'''
from __future__ import annotations
from pathlib import Path
import pandas as pd

#fill and font colors of the flag rows, and of the header
FLAG_COLORS = {'ROCK': ('FFFF00', '000000'), 'RES': ('00AA00', 'FFFFFF'), 'PAY': ('CC0000', 'FFFFFF')}
HEADER_COLORS = ('1F4E79', 'FFFFFF')

#number formats by column name; values keep their full precision
_FORMATS = [(('top', 'bottom', 'base', 'gross', 'net', 'phih', 'hcpv'), '0.00'),
            (('ntg', 'bvw', 'vshale', 'porosity', 'phie', 'saturation', 'sw'), '0.000'),
            (('perm',), '0.0')]


def _number_format(column:str) -> str|None:
    name = column.lower()
    for keys, fmt in _FORMATS:
        if any(k in name for k in keys):
            return fmt
    return None


def _style_sheet(ws, flag_col:str=None):

    '''
    Header, flag colors, borders, column widths and frozen header of a sheet written by pandas
    '''

    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style='thin', color='000000')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    headers = [c.value for c in ws[1]]

    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor=HEADER_COLORS[0])
        cell.font = Font(bold=True, color=HEADER_COLORS[1], size=10)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border
    ws.row_dimensions[1].height = 28

    flag_idx = headers.index(flag_col) if flag_col in headers else None
    formats = [_number_format(str(h)) for h in headers]
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        bg, fg = FLAG_COLORS.get(row[flag_idx].value, ('FFFFFF', '000000')) if flag_idx is not None else ('FFFFFF', '000000')
        for cell, fmt in zip(row, formats):
            cell.fill = PatternFill('solid', fgColor=bg)
            cell.font = Font(color=fg, size=9)
            cell.border = border
            cell.alignment = Alignment(horizontal='center', vertical='center')
            if fmt and isinstance(cell.value, (int, float)):
                cell.number_format = fmt
        ws.row_dimensions[row[0].row].height = 16

    for i, column in enumerate(ws.columns, 1):
        longest = max(len(str(c.value)) if c.value is not None else 0 for c in column)
        ws.column_dimensions[get_column_letter(i)].width = min(max(longest + 2, 8), 22)
    ws.freeze_panes = 'A2'

    if flag_idx is not None:
        #print the summary landscape, one page wide
        ws.page_setup.orientation = 'landscape'
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0


def write_pay_summary(summary:pd.DataFrame, path:Path|str, log_data:pd.DataFrame=None) -> Path:

    r'''
    Writes a pay summary to an Excel workbook: a dark blue header, ROCK rows in yellow, RES in
    green and PAY in red, thin borders, fitted column widths, a frozen header row and rounded
    number formats. The log data, if given, goes to a second sheet

    Parameters
    ----------
    summary : pd.DataFrame
        Pay summary, e.g. `Quanti.paySummary(...).data`, or several wells concatenated. The flag
        column is 'Flag Name' or 'Flag'

    path : pathlib.Path or str
        Workbook to write; '.xlsx' is added if missing

    log_data : pd.DataFrame default None
        Curves and flags, written to a 'Log Data' sheet

    Returns
    -------
    Path of the workbook

    Example
    -------
    >>> from petrolib.io import write_pay_summary
    >>> summaries = [run_well(w).paySummary(w).data for w in wells]
    >>> write_pay_summary(pd.concat(summaries), 'Pay Summary.xlsx')
    '''

    path = Path(path)
    if path.suffix.lower() != '.xlsx':
        path = path.with_name(path.name + '.xlsx')
    flag_col = 'Flag Name' if 'Flag Name' in summary.columns else 'Flag'

    with pd.ExcelWriter(path, engine='openpyxl') as writer:
        summary.to_excel(writer, sheet_name='Pay Summary', index=False)
        _style_sheet(writer.sheets['Pay Summary'], flag_col)
        if log_data is not None:
            log_data.to_excel(writer, sheet_name='Log Data', index=False)
            _style_sheet(writer.sheets['Log Data'])
    return path
